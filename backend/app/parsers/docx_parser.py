import io
from typing import Tuple
import docx
from fastapi import HTTPException

def parse_docx(file_bytes: bytes) -> Tuple[str, bool, bool]:
    """
    Parses DOCX file bytes using python-docx.
    Returns:
        (extracted_text, has_tables, is_single_column)
    """
    if not file_bytes:
        raise HTTPException(status_code=400, detail="Uploaded DOCX file is empty.")
    
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
        
        has_tables = len(doc.tables) > 0
        table_text = []
        if has_tables:
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        table_text.append(row_text)
        
        full_text = "\n".join(paragraphs_text + table_text).strip()
        
        # DOCX column detection heuristic from section properties
        is_single_column = True
        try:
            for section in doc.sections:
                cols = section._sectPr.xpath('./w:cols/@w:num')
                if cols and int(cols[0]) > 1:
                    is_single_column = False
                    break
        except Exception:
            pass
            
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Could not parse DOCX file. Ensure it is a valid document. Error: {str(e)}"
        )
    
    if not full_text:
        raise HTTPException(
            status_code=400,
            detail="No readable text found inside DOCX file."
        )
    
    return full_text, has_tables, is_single_column
